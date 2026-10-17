"""GET /report/duplicates — Excel report of same-amount entries in a date range.

Query: fromDate, toDate (required, YYYY-MM-DD, both days included). Scans expenses,
incomes, car and house expenses (by total) and travel logs (by totalKm); every amount
seen more than once in a table is listed with its entries.
"""

import os
from collections import defaultdict

import boto3
from fiscus.api import api_handler
from fiscus.dates import date_range, within_dates
from fiscus.dynamo import scan_all
from fiscus.xlsx import BOLD, cell_value, workbook_response, write_rows
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

dynamodb = boto3.resource("dynamodb")

TABLES = [
    {
        "name": os.environ["TABLE_EXPENSES"],
        "amount_field": "total",
        "date_field": "invoiceDate",
        "columns": [
            "id",
            "invoiceDate",
            "description",
            "expenseType",
            "total",
            "userId",
            "businessId",
            "notes",
        ],
    },
    {
        "name": os.environ["TABLE_INCOMES"],
        "amount_field": "total",
        "date_field": "invoiceDate",
        "columns": [
            "id",
            "invoiceDate",
            "description",
            "clientName",
            "total",
            "userId",
            "businessId",
            "notes",
        ],
    },
    {
        "name": os.environ["TABLE_CAR_EXPENSES"],
        "amount_field": "total",
        "date_field": "invoiceDate",
        "columns": ["id", "invoiceDate", "description", "expenseType", "total", "carId", "notes"],
    },
    {
        "name": os.environ["TABLE_HOUSE_EXPENSES"],
        "amount_field": "total",
        "date_field": "invoiceDate",
        "columns": ["id", "invoiceDate", "description", "expenseType", "total", "houseId", "notes"],
    },
    {
        "name": os.environ["TABLE_TRAVEL_LOGS"],
        "amount_field": "totalKm",
        "date_field": "travelDate",
        "columns": [
            "id",
            "travelDate",
            "reason",
            "totalKm",
            "userId",
            "carId",
            "businessId",
            "notes",
        ],
    },
]

HEADER_FILL = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
DUPLICATE_FILL = PatternFill(start_color="FFFF99", end_color="FFFF99", fill_type="solid")


def find_duplicates_by_amount(items: list[dict], amount_field: str) -> dict[str, list[dict]]:
    """Group items by amount (as text, e.g. "12.34") and keep groups of two or more."""
    groups: dict[str, list[dict]] = defaultdict(list)
    for item in items:
        amount = item.get(amount_field)
        if amount is not None:
            groups[str(cell_value(amount))].append(item)
    return {amount: group for amount, group in groups.items() if len(group) > 1}


def _write_group(ws, row: int, amount: str, items: list[dict], config: dict) -> int:
    title = ws.cell(row=row, column=1, value=f"Duplicate Amount: {amount} ({len(items)} items)")
    title.font = BOLD
    title.fill = DUPLICATE_FILL
    row += 1
    for index, name in enumerate(config["columns"], start=1):
        cell = ws.cell(row=row, column=index, value=name)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    items.sort(key=lambda item: item.get(config["date_field"], ""))
    return write_rows(ws, row + 1, items, config["columns"]) + 1


@api_handler
def lambda_handler(data, context):
    from_date, to_date = date_range(data)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "Duplicates Report"
    ws.cell(row=1, column=1, value=f"Duplicates Report ({from_date} to {to_date})").font = Font(
        bold=True, size=14
    )
    row = 3
    total_found = 0

    for config in TABLES:
        items = scan_all(
            dynamodb.Table(config["name"]),
            FilterExpression=within_dates(config["date_field"], from_date, to_date),
        )
        duplicates = find_duplicates_by_amount(items, config["amount_field"])
        if not duplicates:
            continue

        ws.cell(row=row, column=1, value=f"Table: {config['name']}").font = Font(bold=True, size=12)
        item_count = sum(len(group) for group in duplicates.values())
        ws.cell(
            row=row + 1,
            column=1,
            value=f"Found {len(duplicates)} duplicate groups ({item_count} items)",
        )
        row += 3
        for amount, group in sorted(duplicates.items(), key=lambda g: float(g[0]), reverse=True):
            row = _write_group(ws, row, amount, group, config)
        total_found += item_count
        row += 1

    summary = (
        f"Total duplicate items found: {total_found}"
        if total_found
        else "No duplicates found in any table."
    )
    ws.cell(row=row, column=1, value=summary).font = BOLD
    return workbook_response(workbook, "DuplicatesReport.xlsx")
