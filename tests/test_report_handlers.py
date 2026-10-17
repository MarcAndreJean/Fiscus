"""Tests for the seven Excel report handlers (moto-backed)."""

import base64
from decimal import Decimal
from io import BytesIO

import pytest
from helpers import load_handler, make_table
from moto import mock_aws
from openpyxl import load_workbook

TABLES = {
    "test-Expenses": "invoiceDate",
    "test-Incomes": "invoiceDate",
    "test-CarExpenses": "invoiceDate",
    "test-HouseExpenses": "invoiceDate",
    "test-TravelLogs": "travelDate",
    "test-MiscFiles": "date",
}


@pytest.fixture
def db():
    with mock_aws():
        yield {name: make_table(name, range_key) for name, range_key in TABLES.items()}


def sheet(result, filename):
    assert result["statusCode"] == 200, result
    assert result["isBase64Encoded"] is True
    assert result["headers"]["Content-Disposition"] == f'attachment; filename="{filename}"'
    return load_workbook(BytesIO(base64.b64decode(result["body"]))).active


def rows(ws):
    return [list(row) for row in ws.iter_rows(values_only=True)]


def row_starting(ws, label):
    matches = [row for row in rows(ws) if label in row[:2]]
    assert matches, f"no row labelled {label!r}"
    return matches[-1]


def amounts(subtotal, tps, tvq, total):
    return {
        k: Decimal(v)
        for k, v in zip(
            ("subtotal", "tps", "tvq", "total"), (subtotal, tps, tvq, total), strict=True
        )
    }


# ---------------------------------------------------------------------------
# Business expenses: sections per expense type, SUM formulas, exact grand total
# ---------------------------------------------------------------------------


def seed_expense(table, item_id, expense_type, date, **money):
    table.put_item(
        Item={
            "id": item_id,
            "userId": "u",
            "businessId": "b",
            "invoiceDate": date,
            "expenseType": expense_type,
            "description": item_id,
            **amounts("0.10", "0", "0", "0.10"),
            **money,
        }
    )


EXPENSE_EVENT = {"userId": "u", "businessId": "b", "fromDate": "2025-01-01", "toDate": "2025-12-31"}


def test_expenses_report_groups_by_type_with_exact_grand_total(db):
    seed_expense(
        db["test-Expenses"], "e1", "advertising", "2025-02-01", **amounts("0.10", "0", "0", "0.10")
    )
    seed_expense(
        db["test-Expenses"], "e2", "utilities", "2025-03-01", **amounts("0.20", "0", "0", "0.20")
    )
    seed_expense(db["test-Expenses"], "e3", "utilities", "2024-12-31")  # outside the range
    handler = load_handler("lambda/report/ReportExpenses.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "ExpensesReport.xlsx")

    assert ws.title == "Expenses Report"
    labels = [row[0] for row in rows(ws)]
    assert "Expense Type: advertising" in labels and "Expense Type: utilities" in labels
    assert "e3" not in labels
    # total is column 14 (N) in the expenses layout
    assert row_starting(ws, "Subtotal utilities")[13].startswith("=SUM(N")
    assert row_starting(ws, "GRAND TOTAL")[13] == 0.3


def test_reports_write_text_that_looks_like_a_formula_as_text(db):
    text = '=HYPERLINK("http://example.com")'
    seed_expense(db["test-Expenses"], text, "utilities", "2025-02-01")
    handler = load_handler("lambda/report/ReportExpenses.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "ExpensesReport.xlsx")

    cells = [cell for row in ws.iter_rows() for cell in row if cell.value == text]
    assert cells
    assert all(cell.data_type == "s" for cell in cells)


def test_reports_drop_control_characters_excel_cannot_store(db):
    # Text pasted from a PDF or Word can carry a form feed or vertical tab.
    seed_expense(db["test-Expenses"], "page 1\x0cpage 2\x0b\x00", "utilities", "2025-02-01")
    handler = load_handler("lambda/report/ReportExpenses.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "ExpensesReport.xlsx")

    assert any("page 1page 2" in row for row in rows(ws))


def test_expenses_report_includes_every_entry_on_to_date(db):
    seed_expense(db["test-Expenses"], "timestamp", "utilities", "2025-12-31T10:00:00.000")
    seed_expense(db["test-Expenses"], "utc", "utilities", "2025-12-31T23:59:59.000Z")
    seed_expense(db["test-Expenses"], "date-only", "utilities", "2025-12-31")
    seed_expense(db["test-Expenses"], "next-day", "utilities", "2026-01-01T00:00:00.000")
    handler = load_handler("lambda/report/ReportExpenses.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "ExpensesReport.xlsx")

    labels = [row[0] for row in rows(ws)]
    assert {"timestamp", "utc", "date-only"} <= set(labels)
    assert "next-day" not in labels


@pytest.mark.parametrize(
    "report,table,date_field,filename",
    [
        ("ReportExpenses", "test-Expenses", "invoiceDate", "ExpensesReport.xlsx"),
        ("ReportIncomes", "test-Incomes", "invoiceDate", "IncomesReport.xlsx"),
        ("ReportCarExpenses", "test-CarExpenses", "invoiceDate", "CarExpensesReport.xlsx"),
        ("ReportHouseExpenses", "test-HouseExpenses", "invoiceDate", "HouseExpensesReport.xlsx"),
        ("ReportMiscFiles", "test-MiscFiles", "date", "MiscFilesReport.xlsx"),
        ("ReportTravelLogs", "test-TravelLogs", "travelDate", "TravelLogsReport.xlsx"),
        ("FindDuplicates", "test-Expenses", "invoiceDate", "DuplicatesReport.xlsx"),
    ],
)
def test_every_report_includes_entries_on_to_date(db, report, table, date_field, filename):
    # Two entries with the same amount so FindDuplicates lists them too.
    for item_id in ("last-day-1", "last-day-2"):
        db[table].put_item(
            Item={
                "id": item_id,
                "userId": "u",
                "businessId": "b",
                "carId": "c",
                "houseId": "h",
                "expenseType": "utilities",
                "category": "others",
                date_field: "2025-12-31T10:00:00.000",
                "totalKm": Decimal("5"),
                "amount": Decimal("5"),
                **amounts("5", "0", "0", "5"),
            }
        )
    handler = load_handler(f"lambda/report/{report}.py")

    ws = sheet(handler.lambda_handler({**EXPENSE_EVENT, "carId": "c"}, None), filename)

    assert {"last-day-1", "last-day-2"} <= {row[0] for row in rows(ws)}


def test_reports_reject_dates_that_are_not_plain_dates(db):
    handler = load_handler("lambda/report/FindDuplicates.py")
    result = handler.lambda_handler({"fromDate": "2025-01-01", "toDate": "31/12/2025"}, None)
    assert result == {"statusCode": 400, "message": "toDate must be a date (YYYY-MM-DD)"}


def test_expenses_report_requires_business(db):
    handler = load_handler("lambda/report/ReportExpenses.py")
    result = handler.lambda_handler({"userId": "u", "fromDate": "a", "toDate": "b"}, None)
    assert result["statusCode"] == 400 and "businessId" in result["message"]


def test_expenses_report_includes_every_page(db):
    padding = "x" * 40_000
    for day in range(1, 31):
        seed_expense(
            db["test-Expenses"], f"r{day:02d}", "utilities", f"2025-04-{day:02d}", notes=padding
        )
    handler = load_handler("lambda/report/ReportExpenses.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "ExpensesReport.xlsx")

    ids = {row[0] for row in rows(ws)}
    assert {f"r{day:02d}" for day in range(1, 31)} <= ids


# ---------------------------------------------------------------------------
# Car and house expenses: type sections, then one block per car/house
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "path,table,owner,owner_label,filename,title",
    [
        (
            "ReportCarExpenses",
            "test-CarExpenses",
            "carId",
            "Car",
            "CarExpensesReport.xlsx",
            "Car Expenses Report",
        ),
        (
            "ReportHouseExpenses",
            "test-HouseExpenses",
            "houseId",
            "House",
            "HouseExpensesReport.xlsx",
            "House Expenses Report",
        ),
    ],
)
def test_owner_report_groups_by_type_then_owner(
    db, path, table, owner, owner_label, filename, title
):
    for item_id, owner_id, total in [("a", "o1", "0.10"), ("b", "o2", "0.20"), ("c", "o1", "1.00")]:
        db[table].put_item(
            Item={
                "id": item_id,
                owner: owner_id,
                "invoiceDate": "2025-05-01",
                "expenseType": "repair",
                **amounts(total, "0", "0", total),
            }
        )
    handler = load_handler(f"lambda/report/{path}.py")

    ws = sheet(
        handler.lambda_handler({"fromDate": "2025-01-01", "toDate": "2025-12-31"}, None), filename
    )

    assert ws.title == title
    assert row_starting(ws, f"{owner_label}: o1") and row_starting(ws, f"{owner_label}: o2")
    # total is column 13 (M) in the car/house layout
    assert row_starting(ws, "Subtotal o1")[12].startswith("=SUM(M")
    assert row_starting(ws, "Total repair")[12] == 1.3
    assert row_starting(ws, "GRAND TOTAL")[12] == 1.3


# ---------------------------------------------------------------------------
# Incomes and travel logs: one flat table with a TOTAL formula row
# ---------------------------------------------------------------------------


def test_incomes_report_is_flat_with_total_formulas(db):
    for item_id, date in [("i2", "2025-06-01"), ("i1", "2025-02-01")]:
        db["test-Incomes"].put_item(
            Item={
                "id": item_id,
                "userId": "u",
                "businessId": "b",
                "invoiceDate": date,
                **amounts("100", "5", "9.98", "114.98"),
            }
        )
    handler = load_handler("lambda/report/ReportIncomes.py")

    ws = sheet(handler.lambda_handler(EXPENSE_EVENT, None), "IncomesReport.xlsx")

    table = rows(ws)
    assert table[0][:3] == ["id", "invoiceDate", "businessId"]
    assert [table[1][0], table[2][0]] == ["i1", "i2"]  # oldest first
    total_row = row_starting(ws, "TOTAL")
    # subtotal/tps/tvq/total are columns 14-17 (N-Q) in the incomes layout
    assert total_row[13:17] == ["=SUM(N2:N3)", "=SUM(O2:O3)", "=SUM(P2:P3)", "=SUM(Q2:Q3)"]


def test_travel_logs_report_sums_km(db):
    for item_id, km in [("t1", "12.5"), ("t2", "7")]:
        db["test-TravelLogs"].put_item(
            Item={
                "id": item_id,
                "userId": "u",
                "businessId": "b",
                "carId": "c",
                "travelDate": "2025-03-0" + item_id[1],
                "totalKm": Decimal(km),
            }
        )
    handler = load_handler("lambda/report/ReportTravelLogs.py")
    event = {**EXPENSE_EVENT, "carId": "c"}

    ws = sheet(handler.lambda_handler(event, None), "TravelLogsReport.xlsx")

    assert ws.title == "Travel Logs Report"
    # totalKm is column 7 (G)
    assert row_starting(ws, "TOTAL")[6] == "=SUM(G2:G3)"
    assert rows(ws)[1][6] == 12.5


def test_travel_logs_report_requires_car(db):
    handler = load_handler("lambda/report/ReportTravelLogs.py")
    result = handler.lambda_handler(EXPENSE_EVENT, None)
    assert result["statusCode"] == 400 and "carId" in result["message"]


# ---------------------------------------------------------------------------
# Misc files: sections per category, amount totals
# ---------------------------------------------------------------------------


def test_misc_files_report_groups_by_category(db):
    for item_id, category, amount in [
        ("m1", "medical", "0.10"),
        ("m2", "medical", "0.20"),
        ("m3", "others", None),
    ]:
        item = {
            "id": item_id,
            "userId": "u",
            "date": "2025-01-0" + item_id[1],
            "category": category,
        }
        if amount:
            item["amount"] = Decimal(amount)
        db["test-MiscFiles"].put_item(Item=item)
    handler = load_handler("lambda/report/ReportMiscFiles.py")

    ws = sheet(
        handler.lambda_handler(
            {"userId": "u", "fromDate": "2025-01-01", "toDate": "2025-12-31"}, None
        ),
        "MiscFilesReport.xlsx",
    )

    assert row_starting(ws, "Category: medical") and row_starting(ws, "Category: others")
    # amount is column 6 (F)
    assert row_starting(ws, "Subtotal medical")[5].startswith("=SUM(F")
    assert row_starting(ws, "GRAND TOTAL")[5] == 0.3


# ---------------------------------------------------------------------------
# Duplicates across tables
# ---------------------------------------------------------------------------


def test_find_duplicates_reports_same_amounts(db):
    seed_expense(
        db["test-Expenses"], "d1", "utilities", "2025-02-01", **amounts("10", "0", "0", "12.34")
    )
    seed_expense(
        db["test-Expenses"], "d2", "utilities", "2025-02-09", **amounts("10", "0", "0", "12.34")
    )
    seed_expense(
        db["test-Expenses"], "d3", "utilities", "2025-02-10", **amounts("10", "0", "0", "99")
    )
    handler = load_handler("lambda/report/FindDuplicates.py")

    ws = sheet(
        handler.lambda_handler({"fromDate": "2025-01-01", "toDate": "2025-12-31"}, None),
        "DuplicatesReport.xlsx",
    )

    labels = [row[0] for row in rows(ws)]
    assert "Duplicate Amount: 12.34 (2 items)" in labels
    assert "d1" in labels and "d2" in labels and "d3" not in labels
    assert "Total duplicate items found: 2" in labels


def test_find_duplicates_with_none(db):
    handler = load_handler("lambda/report/FindDuplicates.py")
    ws = sheet(
        handler.lambda_handler({"fromDate": "2025-01-01", "toDate": "2025-12-31"}, None),
        "DuplicatesReport.xlsx",
    )
    assert "No duplicates found in any table." in [row[0] for row in rows(ws)]


@pytest.mark.parametrize(
    "report",
    [
        "ReportExpenses",
        "ReportIncomes",
        "ReportCarExpenses",
        "ReportHouseExpenses",
        "ReportMiscFiles",
        "ReportTravelLogs",
        "FindDuplicates",
    ],
)
def test_reversed_dates_are_400(db, report):
    handler = load_handler(f"lambda/report/{report}.py")
    event = {**EXPENSE_EVENT, "carId": "c", "fromDate": "2025-12-31", "toDate": "2025-01-01"}
    result = handler.lambda_handler(event, None)
    assert result == {"statusCode": 400, "message": "fromDate must be on or before toDate"}


def test_reports_require_dates(db):
    handler = load_handler("lambda/report/FindDuplicates.py")
    result = handler.lambda_handler({}, None)
    assert result["statusCode"] == 400
    assert "fromDate" in result["message"] and "toDate" in result["message"]
