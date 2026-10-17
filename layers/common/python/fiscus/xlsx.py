"""Excel report building for the report handlers.

Only report functions import this module: it needs openpyxl, which ships in the
OpenpyxlLayer. Three sheet shapes cover every report:

- ``write_flat_table``: header, one row per item, a TOTAL row of SUM formulas.
- ``write_grouped_sections``: a section per group (header, rows, SUM subtotal) and a
  GRAND TOTAL of computed values.
- ``write_owner_sections``: a section per expense type, a block per owner (car/house)
  inside it, type totals and a GRAND TOTAL.

Amounts are summed as ``Decimal`` (how DynamoDB returns them) and converted to float only
when written to a cell.
"""

import base64
from collections import defaultdict
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

XLSX_CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
BOLD = Font(bold=True)
TAXED_AMOUNT_SUMS = ["subtotal", "tps", "tvq", "total"]


def cell_value(value: Any) -> Any:
    return float(value) if isinstance(value, Decimal) else value


def decimal_sum(items: list[dict], field: str) -> Decimal:
    return sum((Decimal(item[field]) for item in items if item.get(field) is not None), Decimal(0))


def sum_columns(columns: list[str], names: list[str]) -> dict[str, int]:
    """1-based column index of each summed field, derived from the column layout."""
    return {name: columns.index(name) + 1 for name in names}


def workbook_response(workbook: Any, filename: str) -> dict:
    buffer = BytesIO()
    workbook.save(buffer)
    return {
        "statusCode": 200,
        "headers": {
            "Content-Type": XLSX_CONTENT_TYPE,
            "Content-Disposition": f'attachment; filename="{filename}"',
        },
        "body": base64.b64encode(buffer.getvalue()).decode("utf-8"),
        "isBase64Encoded": True,
    }


def _label(ws: Any, row: int, column: int, text: str, bold: bool = True) -> None:
    # Labels include stored text (expense type, owner id), so clean them the same way.
    cell = ws.cell(row=row, column=column, value=ILLEGAL_CHARACTERS_RE.sub("", text))
    if bold:
        cell.font = BOLD


def write_header(ws: Any, row: int, columns: list[str], bold: bool = True) -> int:
    for index, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=index, value=name)
        if bold:
            cell.font = BOLD
    return row + 1


def write_value(ws: Any, row: int, column: int, value: Any) -> None:
    """Write one stored value. Text stays text: a note like "=see invoice" is not a formula.

    Control characters Excel cannot store (a form feed pasted from a PDF, for example) are
    dropped; openpyxl would otherwise fail the whole report.
    """
    if isinstance(value, str):
        value = ILLEGAL_CHARACTERS_RE.sub("", value)
    cell = ws.cell(row=row, column=column, value=cell_value(value))
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"


def write_rows(ws: Any, row: int, items: list[dict], columns: list[str]) -> int:
    for item in items:
        for index, name in enumerate(columns, start=1):
            write_value(ws, row, index, item.get(name, ""))
        row += 1
    return row


def _sum_formulas(
    ws: Any, row: int, sums: dict[str, int], first: int, last: int, bold: bool
) -> None:
    for index in sums.values():
        letter = get_column_letter(index)
        cell = ws.cell(row=row, column=index, value=f"=SUM({letter}{first}:{letter}{last})")
        if bold:
            cell.font = BOLD


def _totals(ws: Any, row: int, items: list[dict], sums: dict[str, int]) -> None:
    for name, index in sums.items():
        ws.cell(row=row, column=index, value=cell_value(decimal_sum(items, name))).font = BOLD


def _grouped(items: list[dict], field: str, missing: str, date_field: str) -> dict[str, list]:
    groups: dict[str, list] = defaultdict(list)
    for item in items:
        groups[item.get(field, missing)].append(item)
    for group in groups.values():
        group.sort(key=lambda item: item.get(date_field, ""))
    return dict(sorted(groups.items()))


def write_flat_table(ws: Any, items: list[dict], columns: list[str], sum_names: list[str]) -> None:
    last = write_rows(ws, write_header(ws, 1, columns, bold=False), items, columns)
    ws.cell(row=last, column=1, value="TOTAL")
    if items:
        _sum_formulas(ws, last, sum_columns(columns, sum_names), 2, last - 1, bold=False)


def write_grouped_sections(
    ws: Any,
    items: list[dict],
    *,
    group_field: str,
    group_title: str,
    date_field: str,
    columns: list[str],
    sum_names: list[str],
) -> None:
    sums = sum_columns(columns, sum_names)
    row = 1
    for group, group_items in _grouped(items, group_field, "Uncategorized", date_field).items():
        _label(ws, row, 1, f"{group_title}: {group}")
        first = write_header(ws, row + 1, columns)
        row = write_rows(ws, first, group_items, columns)
        _label(ws, row, 1, f"Subtotal {group}")
        _sum_formulas(ws, row, sums, first, row - 1, bold=True)
        row += 2
    _label(ws, row, 1, "GRAND TOTAL")
    _totals(ws, row, items, sums)


def write_owner_sections(
    ws: Any,
    items: list[dict],
    *,
    owner_field: str,
    owner_title: str,
    columns: list[str],
    sum_names: list[str],
) -> None:
    sums = sum_columns(columns, sum_names)
    row = 1
    for expense_type, type_items in _grouped(items, "expenseType", "Uncategorized", "").items():
        _label(ws, row, 1, f"Expense Type: {expense_type}")
        row += 1
        owners = _grouped(type_items, owner_field, "Unknown", "invoiceDate")
        for owner, owner_items in owners.items():
            _label(ws, row, 2, f"{owner_title}: {owner}")
            first = write_header(ws, row + 1, columns)
            row = write_rows(ws, first, owner_items, columns)
            _label(ws, row, 2, f"Subtotal {owner}")
            _sum_formulas(ws, row, sums, first, row - 1, bold=True)
            row += 2
        _label(ws, row, 1, f"Total {expense_type}")
        _totals(ws, row, type_items, sums)
        row += 2
    _label(ws, row, 1, "GRAND TOTAL")
    _totals(ws, row, items, sums)
